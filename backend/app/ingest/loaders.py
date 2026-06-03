"""Loaders de documents: PDF, DOCX, MD, HTML (i TXT com a cas trivial).

Cada loader retorna una llista de `Pagina` (text + número de pàgina opcional).
Les importacions de biblioteques pesades són diferides perquè importar el mòdul
no exigeixi tenir-les instal·lades (útil per a tests d'altres parts).
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

# Extensions reconegudes → tipus del contracte (pdf|docx|xlsx|md|html).
EXTENSIONS: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".md": "md",
    ".markdown": "md",
    ".html": "html",
    ".htm": "html",
    ".txt": "md",  # text pla tractat com a markdown simple
    # Imatges → OCR (Tesseract). Documents escanejats, fotos de circulars, etc.
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".tiff": "image",
    ".tif": "image",
    ".bmp": "image",
}


@dataclass
class Pagina:
    """Fragment de text d'un document amb el seu número de pàgina (si escau)."""

    text: str
    pagina: int | None = None


def tipus_de(cami: Path) -> str | None:
    """Retorna el tipus del contracte per a un fitxer, o None si no és suportat."""
    return EXTENSIONS.get(cami.suffix.lower())


def _ocr_imatge(imatge) -> str:
    """Aplica OCR (Tesseract, català+castellà) a una imatge PIL. Degrada a '' si falla."""
    try:
        import pytesseract

        return (pytesseract.image_to_string(imatge, lang="cat+spa") or "").strip()
    except Exception:
        return ""


def _ocr_pdf_pagina(cami: Path, pagina: int) -> str:
    """Renderitza una pàgina d'un PDF a imatge i n'extreu el text per OCR."""
    try:
        from pdf2image import convert_from_path

        imatges = convert_from_path(
            str(cami), first_page=pagina, last_page=pagina, dpi=200
        )
        return _ocr_imatge(imatges[0]) if imatges else ""
    except Exception:
        return ""


def _carrega_imatge(cami: Path) -> list[Pagina]:
    """Carrega una imatge i n'extreu el text per OCR."""
    try:
        from PIL import Image

        text = _ocr_imatge(Image.open(cami))
    except Exception:
        text = ""
    return [Pagina(text=text)] if text else []


def _carrega_pdf(cami: Path) -> list[Pagina]:
    from pypdf import PdfReader

    lector = PdfReader(str(cami))
    pagines: list[Pagina] = []
    for i, pag in enumerate(lector.pages, start=1):
        text = (pag.extract_text() or "").strip()
        if not text:
            # PDF escanejat (sense capa de text): OCR de la pàgina.
            text = _ocr_pdf_pagina(cami, i)
        if text:
            pagines.append(Pagina(text=text, pagina=i))
    return pagines


def _carrega_docx(cami: Path) -> list[Pagina]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(cami))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    # Inclou també el text de les taules.
    for taula in doc.tables:
        for fila in taula.rows:
            cel = " | ".join(c.text.strip() for c in fila.cells if c.text.strip())
            if cel:
                parts.append(cel)
    text = "\n".join(parts).strip()
    return [Pagina(text=text)] if text else []


def _carrega_html(cami: Path) -> list[Pagina]:
    from bs4 import BeautifulSoup

    contingut = cami.read_text(encoding="utf-8", errors="ignore")
    sopa = BeautifulSoup(contingut, "html.parser")
    for tag in sopa(["script", "style"]):
        tag.decompose()
    text = sopa.get_text(separator="\n").strip()
    return [Pagina(text=text)] if text else []


def _carrega_text(cami: Path) -> list[Pagina]:
    text = cami.read_text(encoding="utf-8", errors="ignore").strip()
    return [Pagina(text=text)] if text else []


def carrega(cami: Path) -> list[Pagina]:
    """Carrega un document i en retorna les pàgines de text.

    Llança ValueError si el tipus no és suportat.
    """
    tipus = tipus_de(cami)
    if tipus is None:
        raise ValueError(f"Tipus de fitxer no suportat: {cami.suffix}")
    if tipus == "pdf":
        return _carrega_pdf(cami)
    if tipus == "docx":
        return _carrega_docx(cami)
    if tipus == "html":
        return _carrega_html(cami)
    if tipus == "image":
        return _carrega_imatge(cami)
    return _carrega_text(cami)  # md / txt
